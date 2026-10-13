import json

import numpy as np
import openpyxl
import pandas as pd
import pytest

from cr_strat.amfi_debt_trades import (
    MF_TRADE_COLUMNS,
    attach_scheme_categories,
    build_mf_trades,
    classify_deal_type,
    classify_security_type,
    infer_security_type,
    load_amfi_debt_trades,
    mapping_report,
    read_mf_trades,
    to_price_trades,
)
from cr_strat.mf_portfolio import parse_mf_portfolio
from cr_strat.schemes import parse_amfi_navall
from cr_strat.trades import TRADE_COLUMNS, clean_trades, isin_daily, read_trades

HEADER = [
    "Name of the Security", "ISIN", "Type of Security", "Security Rating", "Rating Agency", "Transaction Type",
    "Status of Security", "Fund House", "Scheme Name", "Type of Scheme", "Maturity Date", "Residual Days",
    "Deemed Maturity Date", "Trade Date", "Settlement Date", "Quantity Traded", "Face Value Per Unit",
    "Price at which Traded", "Total Interest Accrued", "Value of the Trade", "Yield at which Traded",
    "Yield at which Valued", "Type of Trade",
]

NAVALL = (
    "Scheme Code;ISIN Div Payout/ ISIN Growth;ISIN Div Reinvestment;Scheme Name;Plan;Option;Net Asset Value;Date\n"
    " \nOpen Ended Schemes(Debt Scheme - Corporate Bond Fund)\n \nEdelweiss Mutual Fund\n \n"
    "100001;INF000A01011;-;Edelweiss Corporate Bond Fund;Direct Plan;Growth;30.1;13-Oct-2026\n"
    " \nAditya Birla Sun Life Mutual Fund\n \n"
    "100002;INF000B01012;-;Aditya Birla Sun Life Corporate Bond Fund;Regular Plan;Growth;101.2;13-Oct-2026\n"
    " \nOpen Ended Schemes(Debt Scheme - Credit Risk Fund)\n \nPGIM India Mutual Fund\n \n"
    "100003;INF000C01013;-;PGIM India Credit Risk Fund;Direct Plan;Growth;15.0;13-Oct-2026\n"
    " \nOpen Ended Schemes(Debt Scheme - Liquid Fund)\n \nHDFC Mutual Fund\n \n"
    "100004;INF000D01014;-;HDFC Liquid Fund;Direct Plan;Growth;4000.0;13-Oct-2026\n"
)


def _row(**kw) -> list:
    keys = {h: h.lower().replace(" ", "_") for h in HEADER}
    return [kw.get(keys[h]) for h in HEADER]


def _save(rows: list[list], path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AMFI_Reports"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _full(**kw) -> list:
    """A row in the post-Sep-2020 schema, every field filled."""
    base = dict(status_of_security="Listed", type_of_scheme="Open ended", maturity_date="29 Sept 2026",
                residual_days=1, deemed_maturity_date="29 Sept 2026", trade_date="28 Sept 2026",
                settlement_date="28 Sept 2026", total_interest_accrued=0, rating_agency="-", security_rating="-")
    return _row(**{**base, **kw})


def _old(**kw) -> list:
    """A row in the pre-Sep-2020 schema: no direction, fund house, type, rating, face, price or traded yield."""
    base = dict(security_rating="-", rating_agency="-", status_of_security="-", maturity_date="14 Jan 2029",
                residual_days=3512, trade_date="04 Jun 2019", settlement_date="06 Jun 2019",
                face_value_per_unit=0, price_at_which_traded=0, total_interest_accrued=0, yield_at_which_traded=0)
    return _row(**{**base, **kw})


@pytest.fixture
def isins(isin):
    return {
        "cp": isin("INE100A14AB"), "ncd": isin("INE200B07AA"), "sdl": isin("IN152018029"),
        "tbill": isin("IN002026X12"), "cd": isin("INE300C16AC"), "repo_coll": isin("IN222024031"),
        "gsec": isin("IN002018045"), "bond": isin("INE031A0866"), "goi_serviced": isin("INE261F08BQ"),
        "cp_old": isin("INE733E14AB"), "deb_old": isin("INE001A07AB"), "tbill_old": isin("IN002018X53"),
        "trep_pseudo": isin("INTREP07061"),
    }


@pytest.fixture
def full_file(tmp_path, isins):
    rows = [
        _full(name_of_the_security="TRP_290926", isin="-", type_of_security="TREPS", transaction_type="Buy",
              status_of_security="Not Applicable", fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund",
              quantity_traded=1500000, face_value_per_unit=100, price_at_which_traded=99.9866,
              value_of_the_trade=149979900, yield_at_which_traded=4.89, yield_at_which_valued=4.89,
              type_of_trade="Others - TREPS"),
        _full(name_of_the_security="Alpha Industries Limited CP (MD 29/09/2026)", isin=isins["cp"],
              type_of_security="Commercial Paper", security_rating="A1+", rating_agency="CRISIL",
              transaction_type="BUY", fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund",
              quantity_traded=300, face_value_per_unit=500000, price_at_which_traded=99.985,
              value_of_the_trade=149977500, yield_at_which_traded=5.18, yield_at_which_valued=5.18,
              type_of_trade="Secondary through market"),
        _full(name_of_the_security="8.15% Beta Finance Ltd NCD", isin=isins["ncd"], type_of_security="Debentures",
              security_rating="CARE AA+", rating_agency="CARE", transaction_type="Sell",
              fund_house="EDEL MUTUAL FUND", scheme_name="Edelweiss Corporate Bond Fund - Regular Plan",
              type_of_scheme="OPEN ENDED", trade_date="25 Sept 2026", quantity_traded=50,
              face_value_per_unit="100000", price_at_which_traded=100.0734, total_interest_accrued=12345.5,
              value_of_the_trade=5016015.5, yield_at_which_traded=8.15, yield_at_which_valued=8.16,
              type_of_trade="OTC"),
        _full(name_of_the_security="8.15% Beta Finance Ltd NCD", isin=isins["ncd"], type_of_security="Corporate Bond",
              security_rating="AA+", rating_agency="CARE", transaction_type="buy", fund_house="PGIM MUTUAL FUND",
              scheme_name="PGIM India Credit Risk Fund", quantity_traded=50, face_value_per_unit=100000,
              price_at_which_traded=100.10, value_of_the_trade=5017345.5, yield_at_which_traded=8.14,
              yield_at_which_valued=8.16, type_of_trade="Secondary Through Off-market"),
        _full(name_of_the_security="SGS MAHARASHTRA 7.10% (25/09/2036)", isin=isins["repo_coll"],
              type_of_security="Reverse Repo", transaction_type="Buy", fund_house="HDFC Mutual Fund",
              scheme_name="HDFC Liquid Fund", quantity_traded=1000000, face_value_per_unit=100,
              price_at_which_traded=96.5, value_of_the_trade=96500000, yield_at_which_traded=5.2,
              yield_at_which_valued=5.2, type_of_trade="Reverse Repo"),
        _full(name_of_the_security="5.10% GS REPO290926", isin="XXX000100001", type_of_security="Corporate Bond Repo",
              transaction_type="Buy", fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund",
              quantity_traded=1, face_value_per_unit=100, price_at_which_traded=100, value_of_the_trade=1e8,
              yield_at_which_traded=5.16, yield_at_which_valued=5.16, type_of_trade="Other- Repo"),
        _full(name_of_the_security="8.28% GUJARAT SDL 2029", isin=isins["sdl"], type_of_security="Government Bonds",
              security_rating="SOV", rating_agency="Sovereign", transaction_type="Buy",
              fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund", quantity_traded=100000,
              face_value_per_unit=100, price_at_which_traded=99.5, value_of_the_trade=9950000,
              yield_at_which_traded=7.2, yield_at_which_valued=7.21, type_of_trade="SECONDARY THROUGH MARKET"),
        _full(name_of_the_security="91 DAY T-BILL 24-DEC-2026", isin=isins["tbill"], type_of_security="Tresury Bills",
              security_rating="Sovereign", rating_agency="N.A.", transaction_type="Sell",
              fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund", quantity_traded=1000000,
              face_value_per_unit=100, price_at_which_traded=98.8, value_of_the_trade=98800000,
              yield_at_which_traded=5.3, yield_at_which_valued=5.3, type_of_trade="Market trade"),
        _full(name_of_the_security="Gamma Bank CD", isin=isins["cd"], type_of_security="Certficate Of Deposit",
              security_rating="A1+", rating_agency="ICRA", transaction_type="Buy", fund_house="HDFC Mutual Fund",
              scheme_name="HDFC Liquid Fund", quantity_traded=1000, face_value_per_unit=500000,
              price_at_which_traded=98.5, value_of_the_trade=492500000, yield_at_which_traded=6.1,
              yield_at_which_valued=6.1, type_of_trade="Primary"),
        _full(name_of_the_security="Alpha Industries Limited CP (MD 28/09/2026)", isin=isins["cp"],
              type_of_security="Commercial Paper", security_rating="0", rating_agency="0", transaction_type="Sell",
              fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund", quantity_traded=100,
              face_value_per_unit=500000, price_at_which_traded=0, value_of_the_trade=5e7,
              yield_at_which_traded=0, yield_at_which_valued=0, type_of_trade="Other- Matured"),
        _full(name_of_the_security="6.68% Pay Fix Rec Mibor (MD 29/09/2031)", type_of_security="Interest Rate Swaps",
              transaction_type="Buy", fund_house="HDFC Mutual Fund", scheme_name="HDFC Liquid Fund",
              quantity_traded=0, face_value_per_unit=0, price_at_which_traded=0, value_of_the_trade=5e8,
              yield_at_which_traded=0, yield_at_which_valued=0, type_of_trade="Interest Rate Swaps"),
    ]
    path = tmp_path / "transaction-report_2026-09-28.xlsx"
    _save(rows, path)
    return path


@pytest.fixture
def old_file(tmp_path, isins):
    absl = "ADITYA BIRLA SUN LIFE BANKING & PSU DEBT FUND"
    rows = [
        _old(name_of_the_security="7.26% GOI - (14-JAN-2029)", isin=isins["gsec"], scheme_name=absl,
             quantity_traded=2000000, value_of_the_trade=201500000, yield_at_which_valued=0.0712,
             type_of_trade="Market Trade"),
        _old(name_of_the_security="8.34% HUDCO - (11-JUL-2022) - SR E 2018", isin=isins["bond"], scheme_name=absl,
             quantity_traded=500000, value_of_the_trade=50500000, yield_at_which_valued=7.85,
             type_of_trade="Interscheme"),
        _old(name_of_the_security="8.24% NABARD - (22-MAR-2029) - GOI SERVICED SERIES",
             isin=isins["goi_serviced"], scheme_name="ADITYA BIRLA SUN LIFE CORPORATE BOND FUND",
             quantity_traded=100, value_of_the_trade=1.0e8, yield_at_which_valued=7.6, type_of_trade="Inter-scheme"),
        _old(name_of_the_security="NTPC LTD - CP (18-JUN-2019)", isin=isins["cp_old"], scheme_name=absl,
             quantity_traded=1000, value_of_the_trade=4.98e8, yield_at_which_valued=0.0702,
             type_of_trade="INTERSCHEME"),
        _old(name_of_the_security="8.70% HDFC LTD - (15-DEC-2020) - SERIES U-008", isin=isins["deb_old"],
             scheme_name="Reliance Short Term Fund", quantity_traded=250, value_of_the_trade=2.5e8,
             yield_at_which_valued=8.3, type_of_trade="Off Market trade"),
        _old(name_of_the_security="91 DAY TBILL 06JUN19", isin=isins["tbill_old"], scheme_name=absl,
             quantity_traded=10, value_of_the_trade=1e6, yield_at_which_valued=6.0, type_of_trade="market trade"),
        _old(name_of_the_security="8.28% Gujarat SDL (MD 20/02/2029)", isin=isins["sdl"], scheme_name=absl,
             quantity_traded=10, value_of_the_trade=1e6, yield_at_which_valued=7.7, type_of_trade="MARKET TRADE"),
        _old(name_of_the_security="TREPS_RED_07.06.2019", isin=isins["trep_pseudo"], scheme_name=absl,
             quantity_traded=10, value_of_the_trade=1e6, yield_at_which_valued=5.7, type_of_trade="Market Trade"),
        _old(name_of_the_security="Loans 1,000,000,000.00 INR 5.60% 07/06/2019", isin="TREPS",
             scheme_name="DSP Liquidity Fund", quantity_traded=-1e9, value_of_the_trade=-999800000,
             yield_at_which_valued=5.6, type_of_trade="Market Trade"),
        _old(name_of_the_security="IDFCFIRSTB_04/07/2019_LEND", isin="INE000000000", scheme_name="Bank ETF",
             quantity_traded=100000, value_of_the_trade=-25000, yield_at_which_valued=0,
             type_of_trade="Off market trade"),
        _old(name_of_the_security="REC_FLOAT & PAY_5.69_07/06/2021_BANK", isin="IRS", scheme_name=absl,
             quantity_traded=1e9, value_of_the_trade=1e9, yield_at_which_valued=5.69, type_of_trade="Off Market trade"),
    ]
    path = tmp_path / "transaction-report_2019-06-06.xlsx"
    _save(rows, path)
    return path


def test_full_schema(full_file, isins):
    everything = load_amfi_debt_trades(full_file, keep_no_isin=True)
    assert list(everything.columns) == list(MF_TRADE_COLUMNS)
    assert len(everything) == 11
    assert everything["direction"].tolist() == [
        "BUY", "BUY", "SELL", "BUY", "BUY", "BUY", "BUY", "SELL", "BUY", "SELL", "BUY"]
    assert everything["direction_available"].all()
    assert everything["deal_type"].tolist() == [
        "TREPS", "MARKET", "MARKET", "OFF_MARKET", "REPO", "REPO", "MARKET", "MARKET", "PRIMARY", "OTHER", "OTHER"]
    assert everything["security_type"].tolist() == [
        "TREPS", "CP", "DEBENTURE", "CORPORATE_BOND", "REPO", "REPO", "SDL", "TBILL", "CD", "CP", "OTHER"]
    assert (everything["security_type_source"] == "reported").all()
    # The reverse repo carries its collateral's ISIN, which is not what was traded.
    repo = everything.iloc[4]
    assert pd.isna(repo["isin"]) and repo["isin_raw"] == isins["repo_coll"]

    df = load_amfi_debt_trades(full_file)
    assert df["isin"].tolist() == [isins[k] for k in ("cp", "ncd", "ncd", "sdl", "tbill", "cd", "cp")]
    ncd = df[df["isin"] == isins["ncd"]].set_index("direction")
    sell = ncd.loc["SELL"]
    assert sell["trade_date"] == pd.Timestamp("2026-09-25")
    assert sell["settlement_date"] == pd.Timestamp("2026-09-28")
    assert sell["maturity_date"] == pd.Timestamp("2026-09-29")
    assert sell["face_traded_inr"] == pytest.approx(5e6)
    assert sell["value_inr"] == pytest.approx(5016015.5)
    assert sell["accrued_interest_inr"] == pytest.approx(12345.5)
    assert (sell["rating_grade"], sell["rating_agency"], sell["rating_notch"]) == ("AA+", "CARE", 2)
    assert ncd.loc["BUY", "rating_agency"] == "CARE"
    assert sell["scheme_type"] == "OPEN"
    assert sell["holder_id"] == "MF|edel|edelweiss corporate bond fund"
    assert sell["listing_status"] == "LISTED"

    sdl = df[df["isin"] == isins["sdl"]].iloc[0]
    assert bool(sdl["sovereign"]) and pd.isna(sdl["rating_grade"])
    assert bool(df.loc[df["isin"] == isins["tbill"], "sovereign"].item())
    matured = df[df["deal_type"] == "OTHER"].iloc[0]
    assert np.isnan(matured["price"]) and np.isnan(matured["yield_traded_pct"])
    assert np.isnan(matured["yield_valued_pct"]) and pd.isna(matured["rating_grade"])


def test_reduced_schema(old_file, isins):
    everything = load_amfi_debt_trades(old_file, keep_no_isin=True)
    assert not everything["direction_available"].any()
    assert everything["direction"].isna().all()
    assert everything["deal_type"].tolist()[:5] == [
        "MARKET", "INTER_SCHEME", "INTER_SCHEME", "INTER_SCHEME", "OFF_MARKET"]
    assert everything["security_type"].tolist() == [
        "GSEC", "CORPORATE_BOND", "CORPORATE_BOND", "CP", "DEBENTURE", "TBILL", "SDL", "TREPS", "TREPS", None,
        "OTHER"]
    assert (everything["security_type_source"].dropna() == "inferred").all()

    df = load_amfi_debt_trades(old_file)
    assert len(df) == 7
    assert isins["trep_pseudo"] not in set(df["isin"])
    assert df["price"].isna().all() and df["yield_traded_pct"].isna().all()
    assert df["face_traded_inr"].isna().all() and df["face_value_per_unit"].eq(0).all()
    assert df["yield_valued_pct"].tolist()[:4] == pytest.approx([7.12, 7.85, 7.6, 7.02])
    assert df["amc"].isna().all()
    assert df["holder_id"].iloc[0] == "MF||aditya birla sun life banking and psu debt fund"
    assert df["trade_date"].iloc[0] == pd.Timestamp("2019-06-04")
    assert df["rating_grade"].isna().all() and not df["sovereign"].any()


@pytest.mark.parametrize("text, expected", [
    ("Secondary through market", "MARKET"), ("SECONDARY - THROUGH MARKET", "MARKET"), ("OTC", "MARKET"),
    ("RFQ Secondary through market", "MARKET"), ("Market Trades", "MARKET"), ("Secondary", "MARKET"),
    ("Off-Market Trade", "OFF_MARKET"), ("Secondary through off market", "OFF_MARKET"),
    ("Inter-scheme", "INTER_SCHEME"), ("INTERSCHEME", "INTER_SCHEME"), ("Inter Scheme Transfer", "INTER_SCHEME"),
    ("Primary", "PRIMARY"), ("PRIMARY THROUGH MARKET", "PRIMARY"), ("Primary Trade", "PRIMARY"),
    ("Buyback", "BUYBACK"), ("Put Option exercised", "BUYBACK"),
    ("TREPS - Lending", "TREPS"), ("Other (TREPS)", "TREPS"), ("Treps\\Reverse Repo", "TREPS"),
    ("Other- TREPS Borrowing", "TREPS"), ("Reverse Repo", "REPO"), ("Others-Reverse Repo-Market Trade", "REPO"),
    ("CROMS Lending", "REPO"), ("Corporate Bond Repo", "REPO"),
    ("Others", "OTHER"), ("Others-IRS", "OTHER"), ("Other- Matured", "OTHER"), ("Interest Rate Swaps", "OTHER"),
    ("Something new", "OTHER"), ("", None), (None, None),
])
def test_classify_deal_type(text, expected):
    assert classify_deal_type(text) == expected


@pytest.mark.parametrize("text, expected", [
    ("TREPS", "TREPS"), ("TREP/CBLO DEPOSITS", "TREPS"), ("Reverse Repo", "REPO"), ("ARCL Repo", "REPO"),
    ("Corporate Debt Repo", "REPO"), ("Commercial Papers", "CP"), ("COMMERCIAL PAPER", "CP"),
    ("Certificate of Deposits", "CD"), ("Certficate Of Deposit", "CD"), ("Treasury Bill", "TBILL"),
    ("Tresury Bills", "TBILL"), ("Cash Management Bill", "CMB"), ("State Development Loan", "SDL"),
    ("GOVT. SECURITY", "GSEC"), ("Government Bonds", "GSEC"), ("Corporate Bonds", "CORPORATE_BOND"),
    ("Debentures", "DEBENTURE"), ("Securitised Debt", "SECURITISED"), ("PTC", "SECURITISED"),
    ("Zero Coupon Bond", "ZCB"), ("AT1 Bonds", "AT1"), ("Perpetual Bond", "AT1"), ("Others - IRS", "OTHER"),
    ("Swaps", "OTHER"), ("Mystery", "OTHER"), ("", None),
])
def test_classify_security_type(text, expected):
    assert classify_security_type(text) == expected


def test_infer_security_type(isin):
    assert infer_security_type(isin("INE095A0807"), "IndusInd Bank 9.50% (Basel III Tier 1 AT1) Perpetual") == "AT1"
    assert infer_security_type(isin("INE095A1600"), "XYZ BANK") == "CD"
    assert infer_security_type(isin("IN002019Y11"), "182 Days - TB - M - 05-12-2019") == "TBILL"
    assert infer_security_type("IDIA00211579", "REPO 07Jun19") == "REPO"
    assert infer_security_type("-", "HDFC Bank FD") == "OTHER"
    assert infer_security_type("", "Unknown paper") is None


def test_mapping_report_flags_unknown_spellings(full_file):
    trades = load_amfi_debt_trades(full_file, keep_no_isin=True)
    odd = trades.iloc[[1]].assign(deal_type_raw="Zebra crossing", deal_type="OTHER")
    report = mapping_report(pd.concat([trades, odd], ignore_index=True))
    assert set(report["field"]) == {"deal_type", "security_type"}
    unmatched = report[~report["rule_matched"]]
    assert unmatched[["field", "raw", "mapped", "rows"]].values.tolist() == [["deal_type", "Zebra crossing", "OTHER", 1]]
    row = report[(report["field"] == "security_type") & (report["raw"] == "Tresury Bills")].iloc[0]
    assert row["mapped"] == "TBILL" and row["rows"] == 1


def test_json_pages(tmp_path, isin):
    bond = isin("INE200B07AA")
    new = {
        "name_of_security": "8.15% Beta Finance Ltd NCD", "isin": bond, "typeofsecurity": "Corporate Bond",
        "security_rating": "AA+", "rating_agency": "ICRA", "transaction_type": "SELL", "statusofsecurity": "LISTED",
        "mf_name": "ADITYA BIRLA SUN LIFE MUTUAL FUND", "scheme_name": "ADITYA BIRLA SUN LIFE CORPORATE BOND FUND",
        "typeof_scheme": "Open Ended", "maturity_date": "2029-06-15T00:00:00.000Z", "residual_days": 900,
        "deemed_maturitydate": "2029-06-15T00:00:00.000Z", "trade_date": "2021-06-14T00:00:00.000Z",
        "settlement_date": "2021-06-15T00:00:00.000Z", "quantity_traded": "50", "face_value_per_unit": "10,00,000.00",
        "priceatwhich_traded": "100.0734", "total_interest_accrued": "1,234.50", "value_of_trade": "50,037,934.50",
        "yieldatwhich_traded": "8.1500", "yield_at_which_valued": "8.1600", "type_of_trade": "Secondary through Market",
    }
    old = {**new, "typeofsecurity": None, "security_rating": None, "rating_agency": None, "transaction_type": None,
           "mf_name": None, "face_value_per_unit": None, "priceatwhich_traded": None, "yieldatwhich_traded": None,
           "yield_at_which_valued": "0.0755", "trade_date": "2019-06-04T00:00:00.000Z",
           "settlement_date": "2019-06-06T00:00:00.000Z", "type_of_trade": "Market Trade"}
    page = tmp_path / "settlement_date_15-Jun-2021_p1_ps5.json"
    page.write_text(json.dumps({"data": [new, old], "meta": {"pagination": {"total": 2}}}))
    df = load_amfi_debt_trades(page)
    first, second = df.iloc[0], df.iloc[1]
    assert first["trade_date"] == pd.Timestamp("2021-06-14")
    assert first["face_traded_inr"] == pytest.approx(5e7)
    assert first["accrued_interest_inr"] == pytest.approx(1234.5)
    assert (first["direction"], first["amc"], first["rating_agency"]) == ("SELL", "ADITYA BIRLA SUN LIFE MUTUAL FUND",
                                                                         "ICRA")
    assert first["security_type"] == "CORPORATE_BOND" and second["security_type"] == "DEBENTURE"
    assert pd.isna(second["direction"]) and second["yield_valued_pct"] == pytest.approx(7.55)
    assert df["source_row"].tolist() == [1, 2]

    holiday = tmp_path / "settlement_date_05-Jun-2019_p1_ps5.json"
    holiday.write_text(json.dumps({"data": [], "meta": {"pagination": {"total": 0}}}))
    empty = load_amfi_debt_trades(holiday, keep_no_isin=True)
    assert empty.empty and list(empty.columns) == list(MF_TRADE_COLUMNS)


def test_missing_header_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Something", "Else"])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(ValueError, match="no AMFI debt-transaction header"):
        load_amfi_debt_trades(tmp_path / "x.xlsx")


def test_attach_categories_and_join_to_holdings(tmp_path, full_file, old_file, isins):
    navall = tmp_path / "NAVAll.txt"
    navall.write_text(NAVALL)
    trades, matches = build_mf_trades([full_file, old_file], navall=navall)
    by_scheme = trades.drop_duplicates("holder_name").set_index("holder_name")

    edel = by_scheme.loc["Edelweiss Corporate Bond Fund - Regular Plan"]
    assert edel["amc"] == "Edelweiss Mutual Fund"
    assert edel["holder_id"] == "MF|edelweiss|edelweiss corporate bond fund"
    assert edel["category"] == "Debt Scheme - Corporate Bond Fund"
    assert by_scheme.loc["PGIM India Credit Risk Fund", "amc"] == "PGIM India Mutual Fund"
    assert by_scheme.loc["PGIM India Credit Risk Fund", "category"] == "Debt Scheme - Credit Risk Fund"
    assert by_scheme.loc["HDFC Liquid Fund", "category"] == "Debt Scheme - Liquid Fund"
    # Old rows have no fund house: the AMC is read from the start of the scheme name.
    old_cbf = by_scheme.loc["ADITYA BIRLA SUN LIFE CORPORATE BOND FUND"]
    assert old_cbf["amc"] == "Aditya Birla Sun Life Mutual Fund"
    assert old_cbf["category"] == "Debt Scheme - Corporate Bond Fund"
    assert pd.isna(by_scheme.loc["Reliance Short Term Fund", "amc"])

    m = matches.set_index("scheme_name")
    assert m.loc["Edelweiss Corporate Bond Fund - Regular Plan", "fund_house_raw"] == "EDEL MUTUAL FUND"
    assert m.loc["Edelweiss Corporate Bond Fund - Regular Plan", "amc_match"] == "variant"
    assert m.loc["ADITYA BIRLA SUN LIFE CORPORATE BOND FUND", "amc_match"] == "from_scheme"
    assert m.loc["Reliance Short Term Fund", "match"] == "none"

    # The same scheme parsed from an AMC portfolio file gets the same holder_id.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Scheme Name: Edelweiss Corporate Bond Fund"])
    ws.append(["Portfolio as on 15-Sep-2026"])
    ws.append(["Name of the Instrument", "ISIN", "Rating", "Quantity", "Market Value (Rs. in Lakhs)"])
    ws.append(["Beta Finance NCD", isins["ncd"], "CARE AA+", 100, 100.0])
    wb.save(tmp_path / "edel.xlsx")
    holdings, _ = parse_mf_portfolio(tmp_path / "edel.xlsx", "Edelweiss Mutual Fund")
    assert holdings["holder_id"].item() == edel["holder_id"]

    overrides = pd.DataFrame([{"amc": "", "scheme_name": "Reliance Short Term Fund",
                               "category": "Debt Scheme - Short Duration Fund"}])
    trades, _ = attach_scheme_categories(load_amfi_debt_trades(old_file), parse_amfi_navall(navall), overrides)
    assert set(trades.loc[trades["holder_name"] == "Reliance Short Term Fund", "category"]) == {
        "Debt Scheme - Short Duration Fund"}


def test_to_price_trades(full_file, old_file, isins, tmp_path):
    trades = pd.concat([load_amfi_debt_trades(full_file), load_amfi_debt_trades(old_file)], ignore_index=True)
    inter = trades.iloc[[1]].assign(deal_type="INTER_SCHEME", yield_traded_pct=8.2)
    trades = pd.concat([trades, inter], ignore_index=True)

    price = to_price_trades(trades)
    assert list(price.columns) == list(TRADE_COLUMNS)
    # Market and off-market trades with a traded yield; not the primary CD, the matured CP or old rows.
    assert price["isin"].tolist() == [isins[k] for k in ("cp", "ncd", "ncd", "sdl", "tbill")]
    assert (price["venue"] == "AMFI_MF").all() and (price["n_trades"] == 1).all()
    assert price["yield_pct"].tolist() == pytest.approx([5.18, 8.15, 8.14, 7.2, 5.3])
    assert price["amount_inr"].iloc[0] == pytest.approx(149977500)
    assert len(to_price_trades(trades, include_inter_scheme=True)) == 6

    price.to_csv(tmp_path / "trades_amfi.csv", index=False, date_format="%Y-%m-%d")
    daily = isin_daily(clean_trades(read_trades(tmp_path / "trades_amfi.csv")))
    ncd = daily[daily["isin"] == isins["ncd"]].set_index("date")
    assert ncd["yield_pct"].to_dict() == {pd.Timestamp("2026-09-25"): 8.15, pd.Timestamp("2026-09-28"): 8.14}
    assert set(daily["venues"]) == {"AMFI_MF"}


def test_read_mf_trades_round_trip(full_file, old_file, tmp_path):
    trades = pd.concat([load_amfi_debt_trades(full_file), load_amfi_debt_trades(old_file)], ignore_index=True)
    path = tmp_path / "mf_trades.csv"
    trades.to_csv(path, index=False, date_format="%Y-%m-%d")
    back = read_mf_trades(path)
    assert len(back) == len(trades)
    assert back["trade_date"].tolist() == trades["trade_date"].tolist()
    assert back["direction_available"].tolist() == trades["direction_available"].tolist()
    assert back["sovereign"].tolist() == trades["sovereign"].tolist()
    assert back["face_traded_inr"].fillna(-1).tolist() == trades["face_traded_inr"].fillna(-1).tolist()
    assert back["direction"].tolist() == trades["direction"].tolist()
    assert len(to_price_trades(back)) == len(to_price_trades(trades))
