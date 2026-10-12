import datetime as dt

import openpyxl
import pandas as pd
import pytest

from cr_strat.mf_portfolio import parse_mf_portfolio
from cr_strat.schemes import attach_categories, normalize_scheme_name, parse_amfi_navall

HEADER = ["Name of the Instrument", "ISIN", "Industry / Rating", "Quantity",
          "Market/Fair Value (Rs. in Lakhs)", "% to Net Assets", "Yield of the instrument"]


@pytest.fixture
def workbook(tmp_path, isin):
    a, b, c = isin("INE001A0700"), isin("INE002B0800"), isin("INE003C0900")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CBF"
    ws.append(["ABC Mutual Fund"])
    ws.append(["Scheme Name : ABC Corporate Bond Fund (An open ended debt scheme investing in AA+ and above)"])
    ws.append(["Portfolio as on September 30, 2026"])
    ws.append([])
    ws.append(HEADER)
    ws.append(["Debt Instruments"])
    ws.append(["7.50% X Finance Ltd NCD", a, "CRISIL AA+", "500", 5012.35, 2.5, 7.9])
    ws.append(["8.10% Y Housing Ltd NCD", b, "ICRA AA", 250, "2,498.10", 1.2, 8.3])
    ws.append(["Sub Total", None, None, None, 7510.45, 3.7])
    ws.append(["TREPS", None, None, None, 100.0, 0.1])
    ws.append(["Grand Total", None, None, None, 7610.45, 100])

    stacked = wb.create_sheet("ALL")
    stacked.append(["ABC Credit Risk Fund"])
    stacked.append(["Monthly portfolio as on", dt.datetime(2026, 9, 15)])
    stacked.append(["Name of Instrument", "ISIN", "Rating", "Quantity", "Market Value (Rs. in Crores)", "% to NAV"])
    stacked.append(["Z Infra NCD", c, "CARE A+", 100, 10.5, 4.0])
    stacked.append(["Grand Total", None, None, None, 10.5, 100])
    stacked.append([])
    stacked.append(["Portfolio of ABC Medium Duration Fund as on 15.09.2026"])
    stacked.append(["Name of Instrument", "ISIN", "Rating", "Quantity", "Market Value (Rs. in Crores)", "% to NAV"])
    stacked.append(["X Finance NCD", a, "CRISIL AA+", 40, 4.1, 3.0])

    notes = wb.create_sheet("Notes")
    notes.append(["Notes to portfolio"])
    path = tmp_path / "abc_portfolio.xlsx"
    wb.save(path)
    return path, (a, b, c)


def test_parse_blocks(workbook):
    path, (a, b, c) = workbook
    holdings, report = parse_mf_portfolio(path, "ABC Mutual Fund")
    assert set(report["status"]) == {"parsed", "skipped"}

    cbf = holdings[holdings["holder_name"] == "ABC Corporate Bond Fund"]
    assert list(cbf["isin"]) == [a, b]
    assert (cbf["as_of_date"] == pd.Timestamp("2026-09-30")).all()
    assert cbf["quantity"].tolist() == [500, 250]
    assert cbf["market_value_inr"].tolist() == pytest.approx([5012.35e5, 2498.10e5])
    assert cbf["rating_disclosed"].tolist() == ["CRISIL AA+", "ICRA AA"]

    crf = holdings[holdings["holder_name"] == "ABC Credit Risk Fund"]
    assert crf["isin"].tolist() == [c]
    assert crf["as_of_date"].item() == pd.Timestamp("2026-09-15")
    assert crf["market_value_inr"].item() == pytest.approx(10.5e7)

    mdf = holdings[holdings["holder_name"] == "ABC Medium Duration Fund"]
    assert mdf["isin"].tolist() == [a]
    assert mdf["as_of_date"].item() == pd.Timestamp("2026-09-15")
    assert holdings["holder_id"].str.startswith("MF|abc|").all()


def test_overrides(workbook):
    path, _ = workbook
    holdings, report = parse_mf_portfolio(path, "ABC Mutual Fund", as_of="2026-10-15", mv_unit="rupee")
    assert (holdings["as_of_date"] == pd.Timestamp("2026-10-15")).all()
    first = holdings.iloc[0]
    assert first["market_value_inr"] == pytest.approx(5012.35)


NAVALL = """Scheme Code;ISIN Div Payout/ ISIN Growth;ISIN Div Reinvestment;Scheme Name;Net Asset Value;Date

Open Ended Schemes(Debt Scheme - Corporate Bond Fund)

ABC Mutual Fund

100001;INF000A01011;-;ABC Corporate Bond Fund - Direct Plan - Growth;30.1;10-Oct-2026
100002;INF000A01029;-;ABC Corporate Bond Fund - Regular Plan - IDCW;12.2;10-Oct-2026

Open Ended Schemes(Debt Scheme - Credit Risk Fund)

ABC Mutual Fund

100003;INF000A01037;-;ABC Credit Risk Fund - Direct Plan - Growth;20.5;10-Oct-2026
"""


def test_scheme_categories(tmp_path, workbook):
    nav = tmp_path / "NAVAll.txt"
    nav.write_text(NAVALL)
    master = parse_amfi_navall(nav)
    assert len(master) == 2
    assert normalize_scheme_name("ABC Corporate Bond Fund - Direct Plan - Growth") == "abc corporate bond fund"

    holdings, _ = parse_mf_portfolio(workbook[0], "ABC Mutual Fund")
    overrides = pd.DataFrame([{"amc": "ABC Mutual Fund", "scheme_name": "ABC Medium Duration Fund",
                               "category": "Debt Scheme - Medium Duration Fund"}])
    tagged, report = attach_categories(holdings, master, overrides)
    by_scheme = dict(zip(report["scheme_name"], report["match"]))
    assert by_scheme == {"ABC Corporate Bond Fund": "exact", "ABC Credit Risk Fund": "exact",
                         "ABC Medium Duration Fund": "override"}
    cats = dict(zip(tagged["holder_name"], tagged["category"]))
    assert cats["ABC Corporate Bond Fund"] == "Debt Scheme - Corporate Bond Fund"
